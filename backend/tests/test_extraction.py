from zipfile import ZipFile

import pytest
from docx import Document
from tokenizers import Tokenizer
from tokenizers.models import WordLevel
from tokenizers.pre_tokenizers import Whitespace
from tokenizers.processors import TemplateProcessing

from rag.embedding import vector_literal
from rag.extraction import ExtractionError, Page, extract_pages, token_chunks


@pytest.fixture
def tokenizer():
    vocab = {"[UNK]": 0, "[CLS]": 1, "[SEP]": 2}
    vocab.update({f"palavra{i}": i + 3 for i in range(500)})
    result = Tokenizer(WordLevel(vocab, unk_token="[UNK]"))
    result.pre_tokenizer = Whitespace()
    result.post_processor = TemplateProcessing(
        single="[CLS] $A [SEP]",
        special_tokens=[("[CLS]", 1), ("[SEP]", 2)],
    )
    return result


def test_token_chunks_preserve_end_and_overlap(tokenizer):
    text = " ".join(f"palavra{i}" for i in range(500))
    chunks = token_chunks([Page(text, 7)], tokenizer)
    assert len(chunks) > 4
    assert chunks[0]["page"] == 7
    assert "palavra499" in chunks[-1]["content"]
    assert "palavra94" in chunks[0]["content"] and "palavra94" in chunks[1]["content"]
    assert all(len(tokenizer.encode(chunk["content"]).ids) <= 128 for chunk in chunks)


def test_chunk_limit_is_enforced(tokenizer):
    with pytest.raises(ExtractionError, match="fragmentos"):
        token_chunks([Page(" ".join(f"palavra{i}" for i in range(500)))], tokenizer, max_chunks=2)


def test_text_and_binary_validation(tmp_path):
    path = tmp_path / "doc"
    path.write_bytes("\ufeffOlá\r\nMundo".encode())
    assert extract_pages(path, ".md") == [Page("Olá\nMundo")]
    path.write_bytes(b"\xff\x00")
    with pytest.raises(ExtractionError, match="UTF-8"):
        extract_pages(path, ".txt")
    path.write_text("   ")
    with pytest.raises(ExtractionError, match="Nenhum texto"):
        extract_pages(path, ".txt")


def test_docx_includes_tables(tmp_path):
    path = tmp_path / "doc.docx"
    doc = Document()
    doc.add_paragraph("Política da empresa")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Reembolso de viagem"
    doc.save(path)
    text = " ".join(page.text for page in extract_pages(path, ".docx"))
    assert "Política" in text and "Reembolso" in text


def test_docx_zip_bomb_guard(tmp_path, monkeypatch):
    import rag.extraction as extraction

    monkeypatch.setattr(extraction, "MAX_DOCX_UNCOMPRESSED", 100)
    path = tmp_path / "large.docx"
    with ZipFile(path, "w") as archive:
        archive.writestr("payload", "x" * 101)
    with pytest.raises(ExtractionError, match="50 MB"):
        extract_pages(path, ".docx")


def test_vector_rejects_invalid_dimensions_and_nan():
    with pytest.raises(ValueError):
        vector_literal([0.0] * 383)
    with pytest.raises(ValueError):
        vector_literal([float("nan")] * 384)
    assert vector_literal([0.0] * 384).startswith("[0,0,")
